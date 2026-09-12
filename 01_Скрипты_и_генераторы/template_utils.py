import os
from typing import Optional, List
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "..", "02_Шаблоны_и_ресурсы")
REPORTS_DIR = os.path.join(BASE_DIR, "..", "03_Готовые_доклады")

IDEAL_DOCX_PATH = os.path.join(TEMPLATES_DIR, "Доклад_ОТЦП_Идеал.docx")
LOGO_PATH = os.path.join(TEMPLATES_DIR, "stankin_logo.png")


def enable_auto_update_fields(doc: Document) -> None:
    """
    Включает автоматическое обновление полей (включая оглавление и номера страниц) при открытии MS Word.
    """
    element = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    element.append(update_fields)


def setup_document_styles(doc: Document) -> None:
    """
    Устанавливает стандартные ГОСТ-поля (левое 3 см, правое 1.5 см, верхнее/нижнее 2 см)
    и регистрирует нумерацию страниц в правом нижнем колонтитуле (со 2-й страницы).
    """
    enable_auto_update_fields(doc)

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        section.different_first_page_header_footer = True
        
        # Очищаем нижний колонтитул первой (титульной) страницы
        first_footer = section.first_page_footer
        for child in list(first_footer._element):
            first_footer._element.remove(child)
        first_footer.add_paragraph()

        # Полностью очищаем основной колонтитул от всех элементов (в т.ч. от sdt блоков)
        footer = section.footer
        for child in list(footer._element):
            footer._element.remove(child)
            
        p_footer = footer.add_paragraph()
        p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        r1 = p_footer.add_run()
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        r1._r.append(fldChar1)

        r2 = p_footer.add_run()
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        r2._r.append(instrText)

        r3 = p_footer.add_run()
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(12)
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        r3._r.append(fldChar2)

        r4 = p_footer.add_run()
        r4.font.name = 'Times New Roman'
        r4.font.size = Pt(12)
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        r4._r.append(fldChar3)


def load_ideal_title_page(topic: str, discipline: Optional[str] = None) -> Document:
    """
    Загружает эталонный файл 'Доклад_ОТЦП_Идеал.docx', сохраняя точную верстку титульного листа.
    Обрезает все элементы шаблона строго после надписи 'МОСКВА', удаляя старый контент.
    """
    if not os.path.exists(IDEAL_DOCX_PATH):
        raise FileNotFoundError(f"Файл эталона не найден по пути: {IDEAL_DOCX_PATH}")

    doc = Document(IDEAL_DOCX_PATH)
    setup_document_styles(doc)

    # Обрезаем все элементы документа СТРОГО ПОСЛЕ надписи "МОСКВА" (после 1-й страницы)
    body = doc.element.body
    children = list(body)
    
    moscow_idx = None
    for i, child in enumerate(children):
        if child.tag.endswith('p'):
            p = Paragraph(child, doc)
            txt_clean = p.text.strip()
            if 'МОСКВА' in txt_clean:
                moscow_idx = i
                break
                
    if moscow_idx is not None:
        for child in children[moscow_idx + 1:]:
            if not child.tag.endswith('sectPr'):
                body.remove(child)

    # Аккуратно заменяем текст темы и дисциплины, сохраняя оригинальное выравнивание и шрифты
    for p in doc.paragraphs:
        txt = p.text.strip()
        if 'По теме' in txt:
            for r in p.runs:
                r.text = ''
            run = p.runs[0] if p.runs else p.add_run()
            run.text = f'По теме «{topic}»'
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif discipline and ('ДИСЦИПЛИНЕ' in txt or 'Основы технологий' in txt):
            if txt.startswith('«') and txt.endswith('»'):
                for r in p.runs:
                    r.text = ''
                run = p.runs[0] if p.runs else p.add_run()
                run.text = f'«{discipline}»'
                run.font.name = 'Times New Roman'
                run.font.size = Pt(13)
                run.font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc


def create_stankin_title_page(
    doc: Document,
    topic: str,
    discipline: str = "Основы технологий цифрового производства",
    student_name: str = "Соколов Вадим Игоревич",
    group: str = "ИДБ-25-12",
    teacher_name: str = "Лобанов О.А."
) -> None:
    """
    Генерирует программный титульный лист по образцу МГТУ СТАНКИН, если эталонный файл отсутствует.
    """
    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.first_line_indent = Cm(0)
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(4)
        run_logo = p_logo.add_run()
        run_logo.add_picture(LOGO_PATH, width=Cm(4.2))

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.first_line_indent = Cm(0)
    p_hdr.paragraph_format.line_spacing = 1.15
    p_hdr.paragraph_format.space_after = Pt(6)
    
    r1 = p_hdr.add_run("МИНОБРНАУКИ РОССИИ\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(10)
    r1.font.bold = True
    
    r2 = p_hdr.add_run("федеральное государственное автономное образовательное учреждение высшего образования\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9)
    r2.font.bold = True
    
    r3 = p_hdr.add_run("«Московский государственный технологический университет «СТАНКИН»\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)
    r3.font.bold = True
    
    r4 = p_hdr.add_run("(ФГАОУ ВО «МГТУ «СТАНКИН»)")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(10)
    r4.font.bold = True

    tbl_inst = doc.add_table(rows=1, cols=2)
    tbl_inst.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_inst.autofit = False
    tbl_inst.columns[0].width = Cm(8.0)
    tbl_inst.columns[1].width = Cm(8.5)
    
    cell_l = tbl_inst.cell(0, 0)
    cell_r = tbl_inst.cell(0, 1)
    
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.first_line_indent = Cm(0)
    p_l.paragraph_format.line_spacing = 1.1
    rl = p_l.add_run("Институт\nинформационных\nтехнологий")
    rl.font.name = "Times New Roman"
    rl.font.size = Pt(10)
    rl.font.bold = True
    rl.font.underline = True
    
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.first_line_indent = Cm(0)
    p_r.paragraph_format.line_spacing = 1.1
    rr = p_r.add_run("Кафедра\nуправления и информатики в\nтехнических системах")
    rr.font.name = "Times New Roman"
    rr.font.size = Pt(10)
    rr.font.bold = True

    p_div = doc.add_paragraph()
    p_div.paragraph_format.first_line_indent = Cm(0)
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(24)

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.first_line_indent = Cm(0)
    p_topic.paragraph_format.line_spacing = 1.2
    p_topic.paragraph_format.space_after = Pt(24)
    
    rt1 = p_topic.add_run("ОТЧЕТ О ПОДГОТОВКЕ\nДОКЛАДА ПО ДИСЦИПЛИНЕ\n")
    rt1.font.name = "Times New Roman"
    rt1.font.size = Pt(12)
    
    rt2 = p_topic.add_run(f"«{discipline}»\n\n")
    rt2.font.name = "Times New Roman"
    rt2.font.size = Pt(13)
    rt2.font.bold = True
    
    rt3 = p_topic.add_run(f"По теме «{topic}»")
    rt3.font.name = "Times New Roman"
    rt3.font.size = Pt(13)
    rt3.font.bold = True

    p_st = doc.add_paragraph()
    p_st.paragraph_format.first_line_indent = Cm(0)
    p_st.paragraph_format.line_spacing = 1.2
    p_st.paragraph_format.space_after = Pt(6)
    
    rs1 = p_st.add_run("СТУДЕНТА  ")
    rs1.font.name = "Times New Roman"
    rs1.font.size = Pt(11)
    
    rs2 = p_st.add_run(" 2 ")
    rs2.font.name = "Times New Roman"
    rs2.font.size = Pt(11)
    rs2.font.underline = True
    
    rs3 = p_st.add_run("  КУРСА      ")
    rs3.font.name = "Times New Roman"
    rs3.font.size = Pt(11)
    
    rs4 = p_st.add_run(" бакалавриата ")
    rs4.font.name = "Times New Roman"
    rs4.font.size = Pt(11)
    rs4.font.underline = True
    
    rs5 = p_st.add_run("      ГРУППЫ  ")
    rs5.font.name = "Times New Roman"
    rs5.font.size = Pt(11)
    
    rs6 = p_st.add_run(f" {group} \n")
    rs6.font.name = "Times New Roman"
    rs6.font.size = Pt(11)
    rs6.font.underline = True
    
    rs7 = p_st.add_run("                           (уровень профессионального образования)\n")
    rs7.font.name = "Times New Roman"
    rs7.font.size = Pt(8)
    rs7.font.italic = True

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.first_line_indent = Cm(0)
    p_name.paragraph_format.space_before = Pt(6)
    p_name.paragraph_format.space_after = Pt(12)
    
    rname = p_name.add_run(student_name)
    rname.font.name = "Times New Roman"
    rname.font.size = Pt(12)
    rname.font.bold = True
    rname.font.underline = True

    p_dir = doc.add_paragraph()
    p_dir.paragraph_format.first_line_indent = Cm(0)
    p_dir.paragraph_format.line_spacing = 1.2
    p_dir.paragraph_format.space_after = Pt(24)
    
    rd1 = p_dir.add_run("Направление:            ")
    rd1.font.name = "Times New Roman"
    rd1.font.size = Pt(11)
    rd1.font.underline = True
    
    rd2 = p_dir.add_run("09.03.03 Прикладная информатика\n")
    rd2.font.name = "Times New Roman"
    rd2.font.size = Pt(11)
    
    rd3 = p_dir.add_run("Профиль подготовки:    ")
    rd3.font.name = "Times New Roman"
    rd3.font.size = Pt(11)
    rd3.font.underline = True
    
    rd4 = p_dir.add_run("Управление данными/ Математическое и компьютерное моделирование процессов и систем")
    rd4.font.name = "Times New Roman"
    rd4.font.size = Pt(11)
    rd4.font.bold = True

    tbl_prep = doc.add_table(rows=2, cols=2)
    tbl_prep.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_prep.autofit = False
    tbl_prep.columns[0].width = Cm(10.0)
    tbl_prep.columns[1].width = Cm(6.5)
    
    cell_p0 = tbl_prep.cell(0, 0)
    p0 = cell_p0.paragraphs[0]
    p0.paragraph_format.first_line_indent = Cm(0)
    rp0 = p0.add_run("Принял:")
    rp0.font.name = "Times New Roman"
    rp0.font.size = Pt(11)
    
    cell_p1 = tbl_prep.cell(1, 0)
    p1 = cell_p1.paragraphs[0]
    p1.paragraph_format.first_line_indent = Cm(0)
    rp1 = p1.add_run("Старший преподаватель")
    rp1.font.name = "Times New Roman"
    rp1.font.size = Pt(11)
    
    cell_p2 = tbl_prep.cell(1, 1)
    p2 = cell_p2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.first_line_indent = Cm(0)
    rp2 = p2.add_run(teacher_name)
    rp2.font.name = "Times New Roman"
    rp2.font.size = Pt(11)

    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.first_line_indent = Cm(0)
    p_city.paragraph_format.space_before = Pt(36)
    p_city.paragraph_format.space_after = Pt(0)
    rcity = p_city.add_run("МОСКВА 2026")
    rcity.font.name = "Times New Roman"
    rcity.font.size = Pt(12)


def add_heading_1(doc: Document, text: str) -> Paragraph:
    """
    Добавляет Заголовок 1 уровня по ГОСТу (16pt, Жирный, по центру, 1.5 интервал).
    """
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    
    pPr = p._p.get_or_add_pPr()
    outlineLvl = OxmlElement('w:outlineLvl')
    outlineLvl.set(qn('w:val'), '0')
    pPr.append(outlineLvl)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_heading_2(doc: Document, text: str) -> Paragraph:
    """
    Добавляет Заголовок 2 уровня по ГОСТу (14pt, Жирный, по ширине с отступом 1.25 см).
    """
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True

    pPr = p._p.get_or_add_pPr()
    outlineLvl = OxmlElement('w:outlineLvl')
    outlineLvl.set(qn('w:val'), '1')
    pPr.append(outlineLvl)

    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_heading_3(doc: Document, text: str) -> Paragraph:
    """
    Добавляет Заголовок 3 уровня по ГОСТу (14pt, Жирный, по ширине с отступом 1.25 см).
    """
    p = doc.add_paragraph(style='Heading 3')
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True

    pPr = p._p.get_or_add_pPr()
    outlineLvl = OxmlElement('w:outlineLvl')
    outlineLvl.set(qn('w:val'), '2')
    pPr.append(outlineLvl)

    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_paragraph(doc: Document, text: str) -> Paragraph:
    """
    Добавляет стандартный абзац текста по ГОСТу (Times New Roman 14pt, 1.5 интервал, абзацный отступ 1.25 см, по ширине).
    """
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_text_paragraphs(doc: Document, raw_text: str) -> None:
    """
    Разбивает составной многострочный текст по переходам строк '\\n' и добавляет каждый как форматированный параграф.
    """
    paragraphs = raw_text.split('\n')
    for p_text in paragraphs:
        cleaned = p_text.strip()
        if not cleaned:
            continue
        add_paragraph(doc, cleaned)


def add_toc(doc: Document, title: str = "Оглавление") -> None:
    """
    Вставляет динамическое поле автоматического оглавления MS Word с заголовком 'Оглавление' по центру.
    """
    p_title = doc.add_paragraph()
    p_title.paragraph_format.first_line_indent = Cm(0)
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(12)
    run_t = p_title.add_run(title)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(16)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0, 0, 0)

    p_toc_field = doc.add_paragraph()
    p_toc_field.paragraph_format.first_line_indent = Cm(0)
    
    r1 = p_toc_field.add_run()._r
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    fldChar1.set(qn('w:dirty'), 'true')
    r1.append(fldChar1)

    r2 = p_toc_field.add_run()._r
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    r2.append(instrText)

    r3 = p_toc_field.add_run()._r
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    r3.append(fldChar2)

    r4 = p_toc_field.add_run()._r
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r4.append(fldChar3)


def add_toc_item(doc: Document, title: str, page_num: str, level: int = 1) -> None:
    """
    Вставляет статическую строку оглавления с точечным заполнителем и номером страницы.
    """
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(3 if level == 1 else 1)
    p.paragraph_format.space_after = Pt(1)
    
    if level == 2:
        p.paragraph_format.left_indent = Cm(0.6)
    elif level == 3:
        p.paragraph_format.left_indent = Cm(1.2)

    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:leader'), 'dot')
    tab.set(qn('w:pos'), '9355')
    tabs.append(tab)
    pPr.append(tabs)

    run_title = p.add_run(title + "\t")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(12)
    if level == 1:
        run_title.font.bold = True

    run_page = p.add_run(page_num)
    run_page.font.name = "Times New Roman"
    run_page.font.size = Pt(12)
    if level == 1:
        run_page.font.bold = True


def format_table_grid(table: Table, col_widths: List[Cm]) -> None:
    """
    Применяет четкий стиль сетки (Table Grid), устанавливает точные ширины колонок,
    поля в ячейках, заголовок с серым фоном и защиту от разрыва строк.
    """
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 1. Форматирование всех строк
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))

        for idx, width in enumerate(col_widths):
            cell = row.cells[idx]
            cell.width = width
            
            # Поля ячеек (padding: top/bottom=6pt, left/right=8pt)
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = OxmlElement('w:tcMar')
            
            topMar = OxmlElement('w:top')
            topMar.set(qn('w:w'), '120')
            topMar.set(qn('w:type'), 'dxa')
            tcMar.append(topMar)
            
            botMar = OxmlElement('w:bottom')
            botMar.set(qn('w:w'), '120')
            botMar.set(qn('w:type'), 'dxa')
            tcMar.append(botMar)
            
            leftMar = OxmlElement('w:left')
            leftMar.set(qn('w:w'), '160')
            leftMar.set(qn('w:type'), 'dxa')
            tcMar.append(leftMar)
            
            rightMar = OxmlElement('w:right')
            rightMar.set(qn('w:w'), '160')
            rightMar.set(qn('w:type'), 'dxa')
            tcMar.append(rightMar)
            
            tcPr.append(tcMar)

    # 2. Заголовок таблицы (повтор на каждой странице и серый фон)
    hdr_row = table.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(OxmlElement('w:tblHeader'))

    for cell in hdr_row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F2F2F2')
        tcPr.append(shd)


def save_report(doc: Document, filename: str) -> str:
    """
    Сохраняет созданный документ .docx в папку '03_Готовые_доклады'.
    При открытом файле в MS Word автоматически сохраняет в альтернативный файл с суффиксом '_новый.docx'.
    """
    os.makedirs(REPORTS_DIR, exist_ok=True)
    target_path = os.path.join(REPORTS_DIR, filename)
    try:
        doc.save(target_path)
        print(f"[+] Успешно сохранен доклад: {target_path}")
        return target_path
    except PermissionError:
        base, ext = os.path.splitext(filename)
        alt_filename = f"{base}_новый{ext}"
        alt_path = os.path.join(REPORTS_DIR, alt_filename)
        doc.save(alt_path)
        print(f"[!] Файл '{filename}' открыт в Word. Сохранено в: {alt_path}")
        return alt_path
