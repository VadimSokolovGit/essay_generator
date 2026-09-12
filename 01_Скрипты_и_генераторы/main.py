import os
import re
import json
import argparse
import requests
from typing import Dict, Any, Optional
from docx import Document
from template_utils import (
    load_ideal_title_page,
    create_stankin_title_page,
    add_heading_1,
    add_heading_2,
    add_text_paragraphs,
    add_toc,
    save_report
)

# ---------------------------------------------------------
# КОНФИГУРАЦИЯ API И НЕЙРОСЕТИ
# ---------------------------------------------------------
DEFAULT_API_URL = "https://openrouter.ai/api/v1/chat/completions"
DEFAULT_MODEL = "google/gemini-2.5-pro"

SYSTEM_PROMPT = """
Ты — строгий академический исследователь и технический писатель.
Пиши академическим, плотным научным стилем. 
Категорически запрещено:
1. Использовать водянистые связки, пафосные эпитеты ("краеугольный камень", "играет ключевую роль", "не просто X, а Y").
2. Использовать тире как замену союзам и связкам предложений.
3. Делать списки-триады ради видимости структуры.
Каждое предложение должно нести конкретный технический, архитектурный или прикладной факт.
Отвечай только содержательным текстом параграфа, без вводных реплик ("Конечно", "Вот текст").
"""


def sanitize_filename(topic: str) -> str:
    """
    Преобразует произвольную тему в безопасное имя файла .docx.
    """
    cleaned = re.sub(r'[\\/*?:"<>|]', '', topic)
    cleaned = cleaned.replace(' ', '_').strip('.')
    if not cleaned:
        cleaned = "Академический_Доклад"
    return f"Доклад_{cleaned[:50]}.docx"


def call_llm(prompt: str, api_key: str, model_name: str = DEFAULT_MODEL) -> str:
    """
    Отправляет запрос к OpenRouter API для генерации уникального научного текста.
    """
    key = api_key or os.getenv("OPENROUTER_API_KEY", "")
    if not key or key == "ВАШ_OPENROUTER_ИЛИ_OPENAI_КЛЮЧ":
        raise ValueError(
            "API ключ не настроен.\n"
            "Установите переменную окружения OPENROUTER_API_KEY или передайте ключ через аргумент --api-key (-k)."
        )

    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json"
    }
    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.3
    }
    
    response = requests.post(DEFAULT_API_URL, headers=headers, json=payload, timeout=120)
    response.raise_for_status()
    
    data = response.json()
    return data["choices"][0]["message"]["content"].strip()


def generate_outline(topic: str, api_key: str) -> Dict[str, Any]:
    """
    Генерирует уникальный структурированный план доклада в формате JSON для ЛЮБОЙ темы.
    """
    prompt = f"""
    Составь подробный план академического доклада по теме: "{topic}".
    Формат ответа — строго валидный JSON без markdown-блоков (без ```json):
    {{
      "intro": "Введение",
      "chapters": [
        {{
          "title": "Глава 1. Название",
          "sections": ["1.1 Название", "1.2 Название", "1.3 Название"]
        }},
        {{
          "title": "Глава 2. Название",
          "sections": ["2.1 Название", "2.2 Название", "2.3 Название"]
        }},
        {{
          "title": "Глава 3. Название",
          "sections": ["3.1 Название", "3.2 Название", "3.3 Название"]
        }},
        {{
          "title": "Глава 4. Название",
          "sections": ["4.1 Название", "4.2 Название", "4.3 Название"]
        }}
      ],
      "conclusion": "Заключение",
      "sources": "Список использованных информационных источников"
    }}
    """
    raw = call_llm(prompt, api_key=api_key)
    cleaned = raw.replace("```json", "").replace("```", "").strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError as e:
        raise ValueError(f"Ошибка парсинга структуры доклада от LLM: {e}\nОтвет LLM: {raw}")


def build_report(
    topic: str,
    discipline: str = "Информационные технологии",
    student_name: str = "Соколов Вадим Игоревич",
    group: str = "ИДБ-25-12",
    teacher_name: str = "Лобанов О.А.",
    output_filename: Optional[str] = None,
    api_key: Optional[str] = None
) -> str:
    """
    Универсальный генератор докладов по ГОСТу на ЛЮБУЮ заданную тему.
    """
    if not output_filename:
        output_filename = sanitize_filename(topic)

    print(f"\n==================================================")
    print(f"[*] Старт генерации доклада на тему: '{topic}'")
    print(f"[*] Дисциплина: '{discipline}' | Студент: '{student_name}'")
    print(f"==================================================\n")

    # 1. Подготовка титульного листа
    try:
        doc = load_ideal_title_page(topic=topic, discipline=discipline)
    except Exception as err:
        print(f"[!] Использование шаблонного титульного листа не удалось ({err}), генерируем программный...")
        doc = Document()
        create_stankin_title_page(
            doc=doc,
            topic=topic,
            discipline=discipline,
            student_name=student_name,
            group=group,
            teacher_name=teacher_name
        )

    doc.add_page_break()

    # 2. Генерация структуры (плана) доклада через LLM
    print("[*] Генерация плана и структуры доклада...")
    outline = generate_outline(topic=topic, api_key=api_key)

    # 3. Динамическое оглавление
    add_toc(doc)
    doc.add_page_break()

    # 4. Введение
    print("[*] Генерация: Введение...")
    add_heading_1(doc, "Введение")
    intro_text = call_llm(
        f"Напиши развернутое научное введение для доклада по теме '{topic}'. "
        f"Обязательно включи: актуальность, объект, предмет, цель работы, задачи и методы исследования. "
        f"Объем: не менее 400-500 слов. Строгий научный стиль.",
        api_key=api_key
    )
    add_text_paragraphs(doc, intro_text)
    doc.add_page_break()

    # 5. Основные главы и параграфы
    for chapter in outline.get("chapters", []):
        ch_title = chapter.get("title", "")
        print(f"[*] Генерация: {ch_title}...")
        add_heading_1(doc, ch_title)

        for sec_title in chapter.get("sections", []):
            print(f"    - {sec_title}...")
            add_heading_2(doc, sec_title)
            sec_text = call_llm(
                f"Напиши детальный технический параграф '{sec_title}' в рамках главы '{ch_title}' "
                f"по теме '{topic}'. Давай глубокие технические подробности, архитектурные особенности, "
                f"структуры данных, алгоритмы и реальные примеры. Без воды и общих вводных фраз. "
                f"Объем: не менее 350-450 слов.",
                api_key=api_key
            )
            add_text_paragraphs(doc, sec_text)

        doc.add_page_break()

    # 6. Заключение
    print("[*] Генерация: Заключение...")
    add_heading_1(doc, "Заключение")
    conclusion_text = call_llm(
        f"Напиши развернутое академическое заключение для доклада '{topic}'. "
        f"Оформи ключевые выводы в виде связных нумерованных положений по результатам работы. Объем: 250-350 слов.",
        api_key=api_key
    )
    add_text_paragraphs(doc, conclusion_text)
    doc.add_page_break()

    # 7. Список литературы
    print("[*] Генерация: Список литературы...")
    add_heading_1(doc, "Список использованных информационных источников")
    sources_text = call_llm(
        f"Составь список из 7 реальных профильных авторитетных источников (монографии, статьи, документация) "
        f"для темы '{topic}'. Оформи строго по ГОСТ Р 7.0.100-2018 (с указанием авторов, издательства, года, страниц/URL).",
        api_key=api_key
    )
    add_text_paragraphs(doc, sources_text)

    # 8. Сохранение итогового документа
    saved_path = save_report(doc, output_filename)
    print(f"\n[+] Доклад успешно сформирован и сохранен по пути:\n    {saved_path}\n")
    return saved_path


def main():
    parser = argparse.ArgumentParser(
        description="Универсальный автоматический генератор академических докладов (.docx) по ГОСТу на ЛЮБУЮ тему."
    )
    parser.add_argument("-t", "--topic", type=str, help="Тема доклада (например: 'Квантовые компьютеры и криптография')")
    parser.add_argument("-d", "--discipline", type=str, default="Информационные технологии", help="Название учебной дисциплины")
    parser.add_argument("-s", "--student", type=str, default="Соколов Вадим Игоревич", help="ФИО студента")
    parser.add_argument("-g", "--group", type=str, default="ИДБ-25-12", help="Группа студента")
    parser.add_argument("--teacher", type=str, default="Лобанов О.А.", help="ФИО преподавателя")
    parser.add_argument("-o", "--output", type=str, help="Имя итогового файла .docx")
    parser.add_argument("-k", "--api-key", type=str, help="OpenRouter API Key")
    parser.add_argument("-i", "--interactive", action="store_true", help="Интерактивный ввод параметров в консоли")

    args = parser.parse_args()

    # Если аргументы не переданы, запускаем интерактивный режим
    if args.interactive or not args.topic:
        print("\n========================================================")
        print(" УНИВЕРСАЛЬНЫЙ ГЕНЕРАТОР АКАДЕМИЧЕСКИХ ДОКЛАДОВ ПО ГОСТ ")
        print("========================================================\n")
        
        topic = input("1. Введите ТЕМУ доклада: ").strip()
        if not topic:
            print("[!] Тема доклада не может быть пустой. Выход.")
            return

        discipline = input("2. Введите ДИСЦИПЛИНУ [Основы технологий цифрового производства]: ").strip()
        if not discipline:
            discipline = "Основы технологий цифрового производства"

        student = input("3. Введите ФИО студента [Соколов Вадим Игоревич]: ").strip()
        if not student:
            student = "Соколов Вадим Игоревич"

        group = input("4. Введите ГРУППУ [ИДБ-25-12]: ").strip()
        if not group:
            group = "ИДБ-25-12"

        teacher = input("5. Введите ФИО преподавателя [Лобанов О.А.]: ").strip()
        if not teacher:
            teacher = "Лобанов О.А."

        api_key = input("6. Введите OpenRouter API Key (нажмите Enter для использования из окружения): ").strip()
        if not api_key:
            api_key = None

        try:
            build_report(
                topic=topic,
                discipline=discipline,
                student_name=student,
                group=group,
                teacher_name=teacher,
                api_key=api_key
            )
        except Exception as err:
            print(f"\n[!] Ошибка при генерации доклада: {err}")
    else:
        try:
            build_report(
                topic=args.topic,
                discipline=args.discipline,
                student_name=args.student,
                group=args.group,
                teacher_name=args.teacher,
                output_filename=args.output,
                api_key=args.api_key
            )
        except Exception as err:
            print(f"\n[!] Ошибка при генерации доклада: {err}")


if __name__ == "__main__":
    main()
